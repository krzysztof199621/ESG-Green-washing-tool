#!/usr/bin/env python3
"""Materiał źródłowy i raport audytu greenwashingu: L'Oréal Polska / Garnier PIL."""

from pathlib import Path

import docx
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
LOGO = ROOT.parent.parent / "ESG Hero Day - Materialy Druk" / "assets" / "esg_logo_white.png"
SHOT = ROOT / "zrodla" / "garnier-pielegnacja-twarzy-2026-10-11.png"


def _set_run(run, size, bold=False, italic=False, color=RGBColor(22, 36, 44), font="Arial"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)


def build_material():
    doc = docx.Document()
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)

    NAVY = RGBColor(8, 24, 31)
    MUTED = RGBColor(90, 107, 119)
    DARK = RGBColor(22, 36, 44)

    def p(text, size=10.5, bold=False, italic=False, color=DARK, before=0, after=6):
        par = doc.add_paragraph()
        par.paragraph_format.space_before = Pt(before)
        par.paragraph_format.space_after = Pt(after)
        par.paragraph_format.line_spacing = 1.12
        _set_run(par.add_run(text), size, bold, italic, color)
        return par

    def h(text):
        par = doc.add_paragraph()
        par.paragraph_format.space_before = Pt(14)
        par.paragraph_format.space_after = Pt(4)
        par.paragraph_format.keep_with_next = True
        _set_run(par.add_run(text), 13, True, False, NAVY)

    p("L'ORÉAL POLSKA / GARNIER – Zestaw treści i deklaracji środowiskowych", 16, True, False, NAVY, after=2)
    p(
        "Zbiór haseł, opisów skali A–E i fragmentów metodyki Product Impact Labelling. "
        "Plik jest zapisem źródłowym do audytu. Nie zawiera oceny prawnej.",
        10.5, False, True, MUTED, after=8,
    )

    rows = [
        ("Podmiot / marka", "L’Oréal Polska Sp. z o.o., marka Garnier. Grupa: L’Oréal."),
        ("Adres z regulaminu witryny", "ul. Grzybowska 62, 00-844 Warszawa. KRS 0000063684. NIP 524-030-33-75."),
        ("Witryny", "https://www.garnier.pl/pielegnacja-twarzy · https://www.garnier.pl/pielegnacja-wlosow · https://www.garnier.pl/regulamin"),
        ("Metodyka grupy", "Product Environmental & Social Impact Labelling Methodologies, 25/10/2021, version 2 (loreal.com)."),
        ("Raport marki", "Garnier Sustainability Progress Report 2020 (loreal.com)."),
        ("Data zebrania", "11 października 2026 r."),
        ("Kontekst zebrania", "Komunikat Prezesa UOKiK z 16 września 2026 r. o zarzucie dotyczącym tej komunikacji. W zbiorze są teksty spółki, nie sentencja decyzji."),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.autofit = False
    for i, (k, v) in enumerate(rows):
        a, b = table.rows[i].cells
        a.width, b.width = Cm(4.4), Cm(13.0)
        pa, pb = a.paragraphs[0], b.paragraphs[0]
        pa.paragraph_format.space_after = Pt(1)
        pb.paragraph_format.space_after = Pt(1)
        _set_run(pa.add_run(k), 9, True)
        _set_run(pb.add_run(v), 9, False)

    h("1. Stan witryny garnier.pl w dniu zebrania")
    p("Tytuł strony https://www.garnier.pl/pielegnacja-twarzy:")
    p("„Pielęgnacja twarzy: Ekologiczne kosmetyki dla każdej cery | Garnier”")
    p("Nagłówek widoczny na stronie:")
    p("„Pielęgnacja twarzy”")
    p("„Wybierz rutynę idealną dla Twojej skóry z produktami do pielęgnacji twarzy od Garnier!”")
    p("Tytuł strony https://www.garnier.pl/pielegnacja-wlosow:")
    p("„Pielęgnacja włosów: Poznaj oceniane produkty do włosów | Garnier”")
    p("Opis meta tej strony:")
    p("„Pielęgnacja włosów z produktami ocenianymi pod względem wpływu na środowisko. Odkryj szampony, maski i odżywki dla każdego rodzaju włosów.”")
    p("Nagłówek widoczny na stronie:")
    p("„Poznaj produkty Garnier przeznaczone do pielęgnacji włosów!”")
    p(
        "W treści widocznej tych dwóch stron w dniu 11 października 2026 r. nie było już akapitu o skali od A do E. "
        "Wynik wyszukiwania dla tych samych adresów nadal zwracał poniższy akapit."
    )

    h("2. Akapit skali A–E przypisany do stron Garnier")
    p("Tekst zwracany przy adresie https://www.garnier.pl/pielegnacja-twarzy (to samo brzmienie przy stronie pielęgnacji włosów):")
    p(
        "„Wybierz rutynę idealną dla Twojej skóry z produktami do pielęgnacji twarzy od Garnier i sprawdź jaki mają wpływ na środowisko. "
        "Każdy produkt jest oceniany w skali od A do E, gdzie A oznacza kosmetyk najlepszy w swojej klasie pod względem wpływu na środowisko i społeczności. "
        "Przeanalizowaliśmy tysiące produktów Grupy L’Oréal, aby stworzyć najlepszą metodologię oznaczeń. "
        "Na ostateczny wynik składa się 14 czynników środowiskowych: od pozyskiwania surowców, poprzez produkcję, transport, aż po korzystanie z produktu końcowego i jego utylizację.”"
    )
    p("Hasła przy ocenach, w brzmieniu przytoczonym w komunikacie Prezesa UOKiK z 16 września 2026 r. jako komunikacja spółki:")
    p("„Wybieraj zrównoważone produkty!”")
    p("„najlepszy w swojej klasie pod względem wpływu na środowisko i społeczności”")
    p(
        "Komunikat podaje, że system nazywany przez grupę „Zielony System Znakowania Produktów” (Product Impact Labelling – PIL) "
        "był na stronie marki Garnier od sierpnia 2021 r. do czerwca 2025 r. oraz że oceny uwzględniały m.in. pozyskanie surowców, "
        "produkcję i transport, użytkowanie i utylizację opakowań, a przy literach były zielone grafiki."
    )

    h("3. Metodyka PIL, wersja 2 z 25 października 2021 r. — fragmenty dosłowne")
    p("Źródło: Product Environmental & Social Impact Labelling Methodologies, 25/10/2021 – VERSION 2.")
    quotes = [
        "„Environmental and Social Impact Labelling was designed to inform consumers on the environmental and social impact of their products, and consequently, empower them to make conscious consumption choices.”",
        "„This labelling is based on the key impact assessments of the Sustainable Product Optimisation Tool (SPOT), the methodology L'Oréal has developed together with independent scientists and experts, and aligned with the European Product Environmental Footprint (PEF) guidelines to scientifically evaluate a product's environmental and social impact.”",
        "„Since 2017, all new or renovated L'Oréal products are evaluated with SPOT. By 2020, 96% of them had an an improved profile.”",
        "„Every time we invent or renovate a product we will improve its environmental or social profile against at least one of the following criteria: the new formula reduces the environmental footprint; the new formula uses raw materials that are sustainably sourced or that or derived from Green Chemistry; the new packaging has an improved environmental profile; the new product a positive social impact.”",
        "„Calculating the product environmental footprint of a L'Oréal product is a way of accurately determining its impact on 14 environmental factors such as greenhouse gas emissions, water scarcity, water acidification and biodiversity.”",
        "„While no cosmetic product categories were selected to be part of the official PEF experiment, a “shadow” group for shampoos following the same approach was initiated, led by the professional association Cosmetics Europe.”",
        "„The methodology is applicable to all L’Oréal cosmetic products.”",
        "„In line with these principles, our beauty product portfolio has been divided into 34 product categories, in order for our consumers to compare products and consider sustainability when making their decision.”",
        "„To classify the various products according to their impact, a linear bounded scale method was selected. Under this approach, the 10% of products with the lowest footprint per user dose would be categorised as Class A and the 10% with the highest footprint would be categorised as Class E. The remaining products were assigned to either category B, C or D.”",
        "„These classes are in line with the recommendation of the French Labelling initiative and compliant with the Ecolabel methodology which ensures that the most environmentally friendly products represent 10% to 20% of the products available on the market.”",
        "„For our scores to be meaningful, products A must be significantly better than products B, which are themselves significantly better than products C. We believe that below 15%, the average difference between products is insufficient to justify a different score.”",
        "Tabela czynników: Climate Change, Water Scarcity, Freshwater Eutrophication, Marine Eutrophication, Water Acidification, Land Transformation, Terrestrial Eutrophication, Particulate Matter, Ionising Radiation, Ozone Depletion, Photochemical Ozone Formation — przy tych pozycjach „PEF COMPLIANCE: Yes”. Freshwater Ecotoxicity, Fossil & Mineral Resources Depletion oraz Toxicity via Environment — „PEF COMPLIANCE: No”, z uzasadnieniem dostosowania baz i metod do kosmetyków.",
        "Przykład oznaczenia na rysunku metodyki (kategoria shampoo): „Overall environmental impact”; „Carbon footprint 84g per usage dose / 60,3g per 10ml”; „Water footprint 6,1 per usage dose / 12,7 per 10ml”; przypis: „Grams of CO2 equivalent (\"shampoo\" category average = 100g)”.",
        "„Manufacturing conditions. Made in a responsible plant. Waste recovery: 100%. Renewable energy: 92%.”",
        "„We consider a manufacturing site to be ‘responsible’ when it continuously reduces its CO2 emissions, water consumption and waste production, contributes to the development of local employment (especially for people with disabilities) and gives access to training, irrespective of whether the site belongs to L’Oréal or a subcontractor.”",
        "„Part of waste recovery (%): Definition: waste that is either reused, recycled, or recovered with energy generation (such as energy used for urban heating).”",
        "„Part of renewable energy (%): Definition: renewable energy is energy that is collected from renewable resources, which are naturally replenished on a human timescale, such as sunlight, wind, rain, tides, waves, and geothermal heat.”",
        "„% of bottle made of recycled material: 74%. Recyclable: yes. Refillable or reusable: no. % of cardboard/paper certified FSC or PEFC: 100%.” Przypisy przy tym przykładzie: „(3) Excluding closing system”; „(4) Recycling instructions may vary locally”.",
        "Definicja recyklowalności w metodyce: „by default we base the assessment on the Ellen MacArthur Foundation definition: ‘A packaging or packaging component is recyclable if its successful post-consumer collection, sorting, and recycling is proven to work in practice and at scale.’ As recycling instructions may vary locally, our labelling may be adapted to local market specificities.”",
        "„Calculation method approved by independent scientific experts and data verified by independent auditor Bureau Veritas Certification.”",
        "„The ingredients and components of this product are sourced from suppliers committed to respect the fundamental principles of the UN on labor standards. 3 suppliers committed to the social inclusion of people from vulnerable communities contributed to making this product.”",
    ]
    for q in quotes:
        p(q, before=2)

    h("4. Garnier Sustainability Progress Report 2020 — fragmenty o oznaczeniu")
    p("„In June 2020, we became the first beauty brand to launch a pioneering ‘product environmental and social impact labelling’ scheme, allowing consumers to make more sustainable choices.”")
    p("„Launched in June 2020 in France, our Product Impact Labelling gives each haircare product a sustainability score to help consumers make more responsible choices.”")
    p(
        "„Based on its overall impact value, every product is classified in comparison to other L’Oréal Group products in the same category, "
        "on a scale from A to E, where A is the most favourable profile for the planet.”"
    )
    p(
        "„Developed in close partnership with independent scientists and experts, this impact assessment methodolgy is unique in the beauty industry. "
        "It is aligned with both the European Product Environmental Footprint standards and the planetary boundaries concept.”"
    )
    p(
        "„The methodology behind this labelling system was co-developed by the L’Oréal Group with 11 international and independent scientists and experts (between 2014 and 2016). "
        "Additionally, the methodology and labelling were peer-reviewed in January 2020 by four independent experts. "
        "The application of the methodology and accuracy of the data are verified by the independent auditor Bureau Veritas, which issued a positive audit report on our calculation methods and the accompanying labelling policy. "
        "Our methodology respects the European Commission’s recommendations on Product Environmental Performance (PEF).”"
    )

    h("5. Identyfikacja usługodawcy witryny")
    p(
        "Regulamin świadczenia usług drogą elektroniczną z 10.05.2018, garnier.pl: "
        "„przez L’Oréal Polska Sp. z o.o. z siedzibą w Warszawie (00-844), przy ulicy Grzybowskiej 62, "
        "wpisaną do rejestru przedsiębiorców przez Sąd Rejonowy dla m.st. Warszawy XIII Wydział Gospodarczy "
        "Krajowego Rejestru Sądowego pod numerem: KRS 0000063684, kapitał zakładowy: 1.886.500,00 zł., NIP 524-030-33-75”."
    )
    p("Infolinia na stronie: 22 33 58 650.")
    p("Stopka: „©2019 Garnier LLC”.")

    out = ROOT / "Material_LOreal_Garnier_PIL.docx"
    doc.save(out)
    return out


def build_report():
    doc = docx.Document()
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(1.7)
        section.right_margin = Cm(1.7)
        section.header_distance = Cm(0.45)
        section.footer_distance = Cm(0.4)

    PRIMARY = RGBColor(8, 24, 31)
    NAVY = RGBColor(8, 24, 31)
    ACCENT = RGBColor(42, 138, 176)
    DANGER = RGBColor(190, 18, 60)
    WARNING = RGBColor(180, 83, 9)
    DARK = RGBColor(22, 36, 44)
    MUTED = RGBColor(90, 107, 119)
    WHITE = RGBColor(255, 255, 255)
    CONTENT = Cm(17.6)

    def shade(cell, hex_color):
        cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'))

    def margins(cell, top=80, bottom=80, left=90, right=90):
        cell._tc.get_or_add_tcPr().append(parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{top}" w:type="dxa"/>'
            f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
            f'<w:left w:w="{left}" w:type="dxa"/>'
            f'<w:right w:w="{right}" w:type="dxa"/>'
            f'</w:tcMar>'
        ))

    def borders(cell, color="E2E8F0", sz="4", left_sz=None, left_color=None):
        lc = left_color or color
        ls = left_sz or sz
        cell._tc.get_or_add_tcPr().append(parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="single" w:sz="{ls}" w:space="0" w:color="{lc}"/>'
            f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tcBorders>'
        ))

    def add_p(text, size=10.5, bold=False, italic=False, color=DARK, before=0, after=8):
        par = doc.add_paragraph()
        par.paragraph_format.space_before = Pt(before)
        par.paragraph_format.space_after = Pt(after)
        par.paragraph_format.line_spacing = 1.12
        _set_run(par.add_run(text), size, bold, italic, color)
        return par

    def add_h(text):
        par = doc.add_paragraph()
        par.paragraph_format.space_before = Pt(16)
        par.paragraph_format.space_after = Pt(6)
        par.paragraph_format.keep_with_next = True
        _set_run(par.add_run(text), 14, True, False, NAVY)
        return par

    def callout(title, lines, bg, border, title_color):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        cell = table.cell(0, 0)
        cell.width = CONTENT
        shade(cell, bg)
        margins(cell, 120, 120, 140, 140)
        borders(cell, border, "6", "24", border)
        par = cell.paragraphs[0]
        par.paragraph_format.space_before = Pt(0)
        par.paragraph_format.space_after = Pt(4)
        _set_run(par.add_run("■ " + title), 10, True, False, title_color)
        for line in lines:
            bp = cell.add_paragraph()
            bp.paragraph_format.space_before = Pt(2)
            bp.paragraph_format.space_after = Pt(2)
            bp.paragraph_format.line_spacing = 1.12
            _set_run(bp.add_run(line), 10, False, False, DARK)
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_before = Pt(0)
        spacer.paragraph_format.space_after = Pt(6)

    def page_field(paragraph):
        run = paragraph.add_run()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = " PAGE "
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.append(begin)
        run._r.append(instr)
        run._r.append(end)
        _set_run(run, 8, False, False, MUTED)

    for section in doc.sections:
        header = section.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        hp.paragraph_format.space_before = Pt(0)
        hp.paragraph_format.space_after = Pt(0)
        htable = header.add_table(1, 2, CONTENT)
        htable.autofit = False
        left, right = htable.rows[0].cells
        left.width, right.width = Cm(8.2), Cm(9.4)
        for cell in (left, right):
            shade(cell, "08181F")
            margins(cell, 60, 60, 120, 120)
            borders(cell, "08181F", "2", "2", "08181F")
        lp = left.paragraphs[0]
        lp.paragraph_format.space_before = lp.paragraph_format.space_after = Pt(0)
        lp.add_run().add_picture(str(LOGO), width=Cm(4.4))
        rp = right.paragraphs[0]
        rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        rp.paragraph_format.space_before = Pt(8)
        rp.paragraph_format.space_after = Pt(0)
        _set_run(rp.add_run("AUDYT RYZYK GREENWASHINGU"), 8, True, False, WHITE)
        rp2 = right.add_paragraph()
        rp2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        rp2.paragraph_format.space_before = Pt(0)
        rp2.paragraph_format.space_after = Pt(0)
        _set_run(rp2.add_run("esginstitute.eu"), 8, False, False, RGBColor(129, 192, 223))
        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.paragraph_format.space_before = Pt(2)
        _set_run(fp.add_run("ESG Institute  ·  Audyt Ryzyk Greenwashingu  ·  L'Oréal Polska  ·  str. "), 8, False, False, MUTED)
        page_field(fp)

    add_p("ESG INSTITUTE  ·  RAPORT AUDYTU REGULACYJNEGO I ŚRODOWISKOWEGO", size=9, bold=True, color=ACCENT, after=2)
    add_p("Audyt Ryzyk Greenwashingu: L'ORÉAL POLSKA", size=20, bold=True, color=PRIMARY, before=2, after=2)
    add_p(
        "Weryfikacja Zielonego Systemu Znakowania Produktów (Product Impact Labelling) marki Garnier na garnier.pl, "
        "metodyki grupy L’Oréal z 25 października 2021 r. (wersja 2) oraz fragmentu Garnier Sustainability Progress Report 2020 "
        "wobec dyrektywy (UE) 2024/825 (ECGT), dyrektywy 2005/29/WE, UPNPR, ustawy o ochronie konkurencji i konsumentów oraz UZNK. "
        "Stan zebrania: 11 października 2026 r.",
        size=10.5, italic=True, color=MUTED, after=8,
    )
    shot = doc.add_paragraph()
    shot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shot.paragraph_format.space_before = Pt(2)
    shot.paragraph_format.space_after = Pt(2)
    shot.add_run().add_picture(str(SHOT), width=Cm(15.4))
    add_p(
        "Zrzut https://www.garnier.pl/pielegnacja-twarzy z 11 października 2026 r. Tytuł dokumentu strony: „Pielęgnacja twarzy: Ekologiczne kosmetyki dla każdej cery | Garnier”.",
        size=8.5, italic=True, color=MUTED, after=10,
    )

    meta = [
        ("Badany podmiot", "L’Oréal Polska Sp. z o.o., ul. Grzybowska 62, 00-844 Warszawa, KRS 0000063684, NIP 524-030-33-75. Marka objęta materiałem: Garnier. System ocen opracowała grupa L’Oréal."),
        ("Przedmiot audytu", "Skala A–E i hasła środowiskowe Garnier na garnier.pl (stosowane od sierpnia 2021 r. do czerwca 2025 r.; w dniu audytu zostały tytuły stron i opis meta), metodyka PIL wersja 2 z 25.10.2021 oraz opis tego oznaczenia w Garnier Sustainability Progress Report 2020. Prezes UOKiK postawił zarzut 16.09.2026. Decyzji i kary w zebranym materiale nie ma."),
        ("Sektor działalności", "Kosmetyki (FMCG): pielęgnacja twarzy i włosów marki Garnier, sprzedaż kierowana do konsumenta w Polsce."),
        ("Reżim prawny", "Dyrektywa (UE) 2024/825 (stosowanie od 27.09.2026), dyrektywa 2005/29/WE, wytyczne UCPD 2021/C 526/01, UPNPR art. 5–6, ustawa o ochronie konkurencji i konsumentów art. 24 i 106, UZNK art. 16. Nowela krajowa (druk Senatu 825) przewiduje wejście w życie 27.09.2027 i nie jest ogłoszona. Wniosek Green Claims nie jest prawem obowiązującym."),
        ("Data audytu", "11 października 2026 r."),
        ("Werdykt operacyjny", "NIE PUBLIKOWAĆ W OBECNEJ FORMIE (KOD CZERWONY / REJECT & RE-NARROW)"),
    ]
    meta_table = doc.add_table(rows=len(meta), cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    widths = (Cm(4.2), Cm(13.4))
    for i, (k, v) in enumerate(meta):
        ck, cv = meta_table.rows[i].cells
        ck.width, cv.width = widths
        shade(ck, "F8FAFC")
        shade(cv, "FFF1F2" if i == 5 else "FFFFFF")
        margins(ck, 50, 50, 80, 80)
        margins(cv, 50, 50, 80, 80)
        borders(ck)
        borders(cv)
        ck.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        pk, pv = ck.paragraphs[0], cv.paragraphs[0]
        pk.paragraph_format.space_before = pk.paragraph_format.space_after = Pt(1)
        pv.paragraph_format.space_before = pv.paragraph_format.space_after = Pt(1)
        _set_run(pk.add_run(k), 8.5, True, False, DARK)
        _set_run(pv.add_run(v), 8.5, i == 5, False, DANGER if i == 5 else DARK)
    add_p("", size=6, after=8)

    callout(
        "WERDYKT: NIE PUBLIKOWAĆ W OBECNEJ FORMIE (KOD CZERWONY / REJECT & RE-NARROW)",
        [
            "Materiał spełnia przesłanki czarnej listy załącznika I do dyrektywy 2005/29/WE w brzmieniu nadanym dyrektywą (UE) 2024/825: pkt 2a (własna skala A–E jako oznakowanie zrównoważonego charakteru, bez systemu certyfikacji strony trzeciej i bez organu publicznego), pkt 4a („ekologiczne kosmetyki”, „zrównoważone produkty”, „sustainability score”) oraz pkt 4b (litera A jako wpływ na środowisko i społeczności oraz „96% improved profile” przy poprawie jednego kryterium).",
            "Osobno litera A jest pozycją w rozkładzie portfela L’Oréal (10% najniższego śladu na dawkę), a tekst dla konsumenta w Polsce nazywa ją kosmetykiem najlepszym w swojej klasie. Metodyka zestawia tę klasę z metodyką Ecolabel i z 10–20% produktów dostępnych na rynku. To rozpowszechnianie informacji prawdziwej w sposób mogący wprowadzać w błąd oraz pominięcie zbioru odniesienia (UPNPR art. 5 ust. 2 pkt 2 i art. 6) i reklama porównawcza bez przesłanek UZNK art. 16 ust. 3.",
            "ECGT stosuje się od 27 września 2026 r. Opóźnienie polskiej noweli tego nie zawiesza. Kanał online nie korzysta z taryfy starych zapasów. Prezes UOKiK postawił już zarzut naruszenia zbiorowych interesów konsumentów (komunikat z 16 września 2026 r.). Kary w tym materiale nie ma. Po potwierdzeniu zarzutu art. 106 ustawy o ochronie konkurencji i konsumentów dopuszcza karę do 10% obrotu.",
        ],
        "FFF1F2", "E11D48", DANGER,
    )

    add_h("1. Analiza Wstępna")
    add_p(
        "L’Oréal Polska Sp. z o.o. prowadzi w Polsce witrynę marki Garnier i sprzedaje kosmetyki do pielęgnacji twarzy, włosów, ciała oraz ochrony przeciwsłonecznej. "
        "Grupa L’Oréal zbudowała własną skalę Product Impact Labelling: litery od A do E przy ogólnym wpływie środowiskowym, ze śladem węglowym i wodnym na dawkę użycia oraz z blokiem warunków wytwarzania, opakowania i wpływu społecznego. "
        "Na garnier.pl skala działała od sierpnia 2021 r. do czerwca 2025 r. W dniu audytu moduł liter zniknął z widocznej treści kategorii, a zostały tytuł „Ekologiczne kosmetyki”, tytuł „oceniane produkty” i opis meta o ocenie wpływu na środowisko."
    )
    add_p(
        "Narracja łączy literę z pełnym cyklem życia, 14 czynnikami, ekspertami, PEF, granicami planetarnymi i weryfikacją Bureau Veritas. "
        "Metodyka wersji 2 z 25 października 2021 r. opisuje literę inaczej: klasa A to 10% produktów L’Oréal o najniższym śladzie na dawkę w kategorii portfela grupy, klasa E to 10% o śladzie najwyższym, a B, C i D są przedziałami między tymi ogonami. "
        "Portfel podzielono na 34 kategorie. Raport postępu Garnier z 2020 r. podaje wprost porównanie z innymi produktami grupy L’Oréal. "
        "Akapit kierowany do polskiego konsumenta mówi o „najlepszym w swojej klasie” i o tym, że tysiące produktów grupy posłużyły do zbudowania metodologii."
    )
    add_p(
        "To ten sam układ, który Prezes UOKiK już kwestionuje przy własnych znacznikach i hasłach ogólnych: listek „eko” i „zrównoważone” u Tchibo, „Join Life” u Zary, ogólne „eko” przy braku skali u kurierów. "
        "Tu doszedł element liczbowy. Litera wygląda jak próg jakości środowiskowej, a jest pozycją w rozkładzie własnego portfela. "
        "Część czynników w tabeli metodyki ma zgodność z PEF oznaczoną jako „No”. Kategoria kosmetyczna nie weszła do oficjalnego eksperymentu PEF. "
        "Audyt dotyczy tego przekazu. Nie rozstrzyga postępowania administracyjnego i nie stwierdza nałożenia kary."
    )

    add_h("2. Analiza szczegółowa - tabela zbiorcza naruszeń")
    findings = [
        ("01", "Litera A: „najlepszy w swojej klasie” pod względem środowiska i społeczności", "UPNPR art. 5 ust. 2 pkt 2 i art. 6 · ECGT zał. I pkt 4b · UZNK art. 16 ust. 3", "BARDZO WYSOKIE", "Klasa A to 10% portfela L’Oréal, a komunikat buduje obraz progu rynkowego i skutku społecznego"),
        ("02", "Własna skala A–E jako oznakowanie zrównoważonego charakteru", "ECGT zał. I pkt 2a · UPNPR art. 5 ust. 3 pkt 2", "WYSOKIE", "Znak grupy, bez systemu certyfikacji strony trzeciej i bez organu publicznego"),
        ("03", "„Ekologiczne kosmetyki”, „zrównoważone produkty”, „sustainability score”", "ECGT zał. I pkt 4a · UPNPR art. 5", "WYSOKIE", "Hasło ogólne bez EU Ecolabel i bez ISO 14024 typu I w tym samym kanale"),
        ("04", "Skala „compliant with the Ecolabel methodology” i 10–20% rynku", "UPNPR art. 5 ust. 2 pkt 2 · ECGT zał. I pkt 4a", "WYSOKIE", "Powiązanie własnego rozkładu portfela z metodyką Ecolabel i z udziałem w rynku"),
        ("05", "PEF, 14 czynników, eksperci i Bureau Veritas przy literze", "UPNPR art. 5 ust. 2 pkt 2 i art. 6", "WYSOKIE", "Weryfikacja metody podana jak obiektywny próg; trzy czynniki poza zgodnością PEF"),
        ("06", "„Made in a responsible plant” i „Renewable energy: 92%”", "ECGT zał. I pkt 4a · UPNPR art. 6", "WYSOKIE", "„Responsible” jest hasłem ogólnym; procent energii nie ma granicy systemu przy liczbie"),
        ("07", "„Waste recovery: 100%” przy definicji obejmującej odzysk energii", "UPNPR art. 5 ust. 2 pkt 2 i art. 6", "ŚREDNIE", "100% obejmuje spalanie z odzyskiem energii, a linia oznaczenia tego nie rozdziela"),
        ("08", "„Recyclable: yes” i 74% recyklatu z wyłączeniem zamknięcia", "ECGT zał. I pkt 4b · UPNPR art. 6", "ŚREDNIE", "Tak/nie i procent opakowania przy wyłączeniu zamknięcia i lokalnych zasad zbiórki"),
        ("09", "84 g CO2e na dawkę wobec średniej kategorii 100 g", "UPNPR art. 6 · ISO 14040 / 14044", "ŚREDNIE", "Liczba ma metodykę w PDF, a średnia jest średnią wewnętrznej kategorii"),
        ("10", "Wpływ społeczny: standardy pracy ONZ i „3 suppliers”", "UPNPR art. 5 ust. 3 pkt 2 i art. 6 · ECGT zał. I pkt 4b", "ŚREDNIE", "Atrybut dostawców przeniesiony na literę „wpływu na społeczności” produktu"),
        ("11", "„96% of them had an improved profile”", "ECGT zał. I pkt 4b · UPNPR art. 5 ust. 2 pkt 2", "ŚREDNIE", "Poprawa oznacza co najmniej jedno z kilku kryteriów, także społecznych"),
    ]
    sum_table = doc.add_table(rows=1 + len(findings), cols=5)
    sum_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sum_table.autofit = False
    sum_table.style = "Table Grid"
    col_w = (Cm(1.1), Cm(5.5), Cm(4.3), Cm(2.3), Cm(4.4))
    headers = ["ID", "Analizowane twierdzenie", "Podstawa prawna", "Ryzyko", "Istota naruszenia"]
    for c, htxt in enumerate(headers):
        cell = sum_table.rows[0].cells[c]
        cell.width = col_w[c]
        shade(cell, "08181F")
        margins(cell, 50, 50, 50, 50)
        par = cell.paragraphs[0]
        par.paragraph_format.space_before = par.paragraph_format.space_after = Pt(1)
        _set_run(par.add_run(htxt), 8, True, False, WHITE)
    for r, row in enumerate(findings):
        bg = "FFFFFF" if r % 2 == 0 else "F8FAFC"
        for c, val in enumerate(row):
            cell = sum_table.rows[r + 1].cells[c]
            cell.width = col_w[c]
            shade(cell, bg)
            margins(cell, 40, 40, 50, 50)
            borders(cell)
            par = cell.paragraphs[0]
            par.paragraph_format.space_before = par.paragraph_format.space_after = Pt(1)
            run = par.add_run(val)
            if c == 3:
                color = WARNING if val == "ŚREDNIE" else DANGER
                _set_run(run, 7.5, True, False, color)
            elif c == 0:
                _set_run(run, 8, True, False, NAVY)
            else:
                _set_run(run, 7.5, False, False, DARK)

    add_p("", size=4, after=4)
    add_h("3. Szczegółowa analiza przypadków")

    cases = [
        {
            "title": "PRZYPADEK 01: Litera A jako „najlepszy w klasie” dla środowiska i społeczności",
            "basis": "Podstawa prawna: UPNPR art. 5 ust. 2 pkt 2 i art. 6 · ECGT, załącznik I pkt 4b · UZNK art. 16 ust. 1 pkt 2 i ust. 3 pkt 1–3 · UPNPR art. 5 ust. 2 pkt 3",
            "risk": "Poziom ryzyka: BARDZO WYSOKIE (czarna lista i porównanie)",
            "quotes": [
                "„Każdy produkt jest oceniany w skali od A do E, gdzie A oznacza kosmetyk najlepszy w swojej klasie pod względem wpływu na środowisko i społeczności.”",
                "„Przeanalizowaliśmy tysiące produktów Grupy L’Oréal, aby stworzyć najlepszą metodologię oznaczeń.”",
                "„the 10% of products with the lowest footprint per user dose would be categorised as Class A and the 10% with the highest footprint would be categorised as Class E.”",
                "„every product is classified in comparison to other L’Oréal Group products in the same category, on a scale from A to E, where A is the most favourable profile for the planet.”",
            ],
            "assessment": (
                "Metodyka ustala klasę A jako 10% produktów o najniższym śladzie na dawkę użycia, a klasę E jako 10% o śladzie najwyższym. "
                "Zakres metodyki to produkty kosmetyczne L’Oréal. Portfel grupy podzielono na 34 kategorie. Raport Garnier z 2020 r. nazywa tę operację porównaniem z innymi produktami grupy w tej samej kategorii. "
                "Akapit zwracany przy adresach kategorii garnier.pl oraz hasło przytoczone w komunikacie Prezesa UOKiK definiują literę A jako kosmetyk najlepszy w swojej klasie pod względem wpływu na środowisko i społeczności. "
                "Zdanie o tysiącach produktów grupy dotyczy zbudowania metodologii. Zbiór, wobec którego liczony jest percentyl, w tym akapicie nie pada. "
                "To informacja o istnieniu skali i o pracach nad metodą, podana tak, że litera czyta się jak próg jakości (UPNPR art. 5 ust. 2 pkt 2). "
                "Pominięcie zbioru odniesienia zmienia znaczenie litery (art. 6). "
                "Litera jest śladem środowiskowym. Blok społeczny w metodyce jest osobny: zasady pracy dostawców i liczba dostawców włączających osoby z grup wrażliwych. "
                "Jedna litera „środowiska i społeczności” obejmuje zatem aspekt, którego ta litera nie mierzy (ECGT pkt 4b). "
                "Zwrot „w swojej klasie” pozwala pośrednio rozpoznać kosmetyki tej samej funkcji, w tym spoza grupy. "
                "Reklama porównawcza wymaga łącznie braku wprowadzenia w błąd oraz rzetelnego, sprawdzalnego porównania towarów zaspokajających tę samą potrzebę (UZNK art. 16 ust. 3 pkt 1–3). "
                "Pozycja w portfelu własnym tego warunku nie spełnia. „Most favourable profile for the planet” w raporcie angielskim dokłada do porównania wewnętrznego obraz profilu najkorzystniejszego dla planety."
            ),
            "remediation": (
                "Wycofać literę A–E z kart produktu, kategorii, tytułów, meta, marketplace i materiałów archiwalnych, które nadal ją pokazują. "
                "Zdjąć „najlepszy w swojej klasie”, „most favourable profile for the planet” oraz łączenie litery ze skutkiem społecznym. "
                "Jeżeli zostaje liczba śladu, w tym samym bloku co liczba podać: zbiór odniesienia (produkty L’Oréal w nazwanej kategorii), jednostkę (dawka użycia), granice cyklu życia i datę metodyki."
            ),
            "claim_title": "SFORMUŁOWANIE NIE DO URATOWANIA",
            "claim": (
                "Litera A jako „najlepszy w swojej klasie” pod względem środowiska i społeczności nie ma bezpiecznego brzmienia. "
                "Dopisek „wśród produktów L’Oréal” nie zamienia własnej skali literowej w oznakowanie zrównoważonego charakteru. Skalę literową usunąć. "
                "Informacja liczbowa, jeżeli spółka zdecyduje się ją zostawić, ma formę z przypadku 09."
            ),
            "salvage": False,
        },
        {
            "title": "PRZYPADEK 02: Własna skala A–E jako oznakowanie zrównoważonego charakteru",
            "basis": "Podstawa prawna: ECGT, załącznik I pkt 2a · UPNPR art. 5 ust. 3 pkt 2",
            "risk": "Poziom ryzyka: WYSOKIE (czarna lista)",
            "quotes": [
                "„Environmental and Social Impact Labelling was designed to inform consumers on the environmental and social impact of their products.”",
                "„we became the first beauty brand to launch a pioneering ‘product environmental and social impact labelling’ scheme.”",
                "„Calculation method approved by independent scientific experts and data verified by independent auditor Bureau Veritas Certification.”",
            ],
            "assessment": (
                "Punkt 2a zakazuje wyświetlania oznakowania dotyczącego zrównoważonego charakteru, które nie opiera się na systemie certyfikacji ani nie zostało ustanowione przez organ publiczny. "
                "Oznakowanie w ECGT to dobrowolny znak zaufania lub jakości odnoszący się do cech środowiskowych lub społecznych. "
                "System certyfikacji wymaga weryfikacji przez niezależną stronę trzecią, warunków dostępnych publicznie, otwartości, procedury niezgodności i monitoringu niezależnego od właściciela systemu oraz od przedsiębiorcy. "
                "Skala A–E jest znakiem grupy L’Oréal. Bureau Veritas, według tekstu spółki, weryfikuje metody obliczeń, dane i politykę oznaczania. "
                "Recenzja ekspertów dotyczy metody. Żaden z tych opisów nie jest certyfikatem produktu w systemie otwartym dla innych przedsiębiorców i niezależnym od właściciela skali. "
                "EU Ecolabel i urzędowo uznany ISO 14024 typu I są w bazie wiedzy wzorcem uznanej wysokiej efektywności. Tego znaku przy literze nie ma. "
                "Zielona prezentacja ocen, opisana w komunikacie UOKiK, wzmacnia odbiór litery jako znaku zaufania. FAQ Komisji z 30 czerwca 2026 r. traktuje zieleń przy haśle jako możliwy znak zrównoważonego charakteru."
            ),
            "remediation": (
                "Zaprzestać eksponowania litery A–E, zielonego znaczka oceny i nazwy „Zielony System Znakowania Produktów” jako znaku przy produkcie. "
                "Weryfikacji Bureau Veritas nie przedstawiać w miejscu, w którym konsument czyta ją jak certyfikat kosmetyku."
            ),
            "claim_title": "SFORMUŁOWANIE NIE DO URATOWANIA",
            "claim": (
                "Własna skala literowa nie przechodzi w dozwolony znak przez dopisek o audytorze lub o ekspertach. "
                "Znak przy produkcie może być EU Ecolabel albo urzędowo uznanym oznakowaniem ISO 14024 typu I, jeżeli produkt ten znak ma i hasło mieści się w jego zakresie. "
                "W zebranym materiale takiego znaku nie podano."
            ),
            "salvage": False,
        },
        {
            "title": "PRZYPADEK 03: Ogólne hasła ekologiczne bez uznanej wysokiej efektywności",
            "basis": "Podstawa prawna: ECGT, załącznik I pkt 4a · UPNPR art. 5",
            "risk": "Poziom ryzyka: WYSOKIE (czarna lista)",
            "quotes": [
                "„Pielęgnacja twarzy: Ekologiczne kosmetyki dla każdej cery | Garnier”",
                "„Wybieraj zrównoważone produkty!”",
                "„Pielęgnacja włosów z produktami ocenianymi pod względem wpływu na środowisko.”",
                "„a sustainability score to help consumers make more responsible choices” oraz „allowing consumers to make more sustainable choices.”",
            ],
            "assessment": (
                "Punkt 4a zakazuje ogólnego twierdzenia dotyczącego ekologiczności, gdy przedsiębiorca nie wykazuje uznanej wysokiej efektywności ekologicznej istotnej dla tego twierdzenia. "
                "Uznana efektywność to EU Ecolabel (rozporządzenie 66/2010), urzędowo uznany program ISO 14024 typu I albo najwyższy poziom z innego prawa Unii, i musi odpowiadać całemu twierdzeniu. "
                "„Ekologiczne”, „zrównoważone”, „sustainability” i „more responsible choices” są hasłami ogólnymi. W materiale nie ma EU Ecolabel ani typu I na kosmetyk Garnier. "
                "Specyfikacja zdejmuje status hasła ogólnego tylko w tym samym kanale i tylko gdy dotyczy tego hasła. Tytuł strony i meta są pierwszym komunikatem wyniku wyszukiwania. "
                "PDF metodyki na innej witrynie ich nie dookreśla. CPC (czerwiec 2026): od 27 września 2026 r. kanał online nie ma taryfy starych zapasów. "
                "Tytuł „Ekologiczne kosmetyki” i meta o ocenie wpływu były dostępne 11 października 2026 r., już po dacie stosowania ECGT i po zaprzestaniu modułu liter w czerwcu 2025 r. "
                "Wzorce UOKiK w bazie: „eko” i „zrównoważone” u Tchibo oraz ogólne hasła „Join Life” u Zary."
            ),
            "remediation": (
                "Zmienić title i meta garnier.pl, w tym „Ekologiczne kosmetyki”, „oceniane produkty” i „wpływ na środowisko” bez zakresu. "
                "Zdjąć „Wybieraj zrównoważone produkty”, „sustainability score”, „more sustainable choices” i „more responsible choices” z materiałów, które nadal cytują skalę. "
                "Zlecić odświeżenie snippetów wyszukiwarek po zmianie title i meta."
            ),
            "claim_title": "DOZWOLONE SFORMUŁOWANIE (WZÓR PO KOREKCIE)",
            "claim": (
                "„Pielęgnacja twarzy Garnier. Przy wybranym kosmetyku podajemy ślad węglowy jednej dawki użycia, z granicami cyklu życia i datą metodyki, albo nie podajemy twierdzenia środowiskowego.” "
                "Bez słów „ekologiczny”, „zrównoważony” i „sustainability score”."
            ),
            "salvage": True,
        },
        {
            "title": "PRZYPADEK 04: Powiązanie własnej skali z metodyką Ecolabel i z udziałem w rynku",
            "basis": "Podstawa prawna: UPNPR art. 5 ust. 2 pkt 2 · ECGT, załącznik I pkt 4a",
            "risk": "Poziom ryzyka: WYSOKIE (czarna lista przy haśle ogólnej efektywności)",
            "quotes": [
                "„These classes are in line with the recommendation of the French Labelling initiative and compliant with the Ecolabel methodology which ensures that the most environmentally friendly products represent 10% to 20% of the products available on the market.”",
                "„the 10% of products with the lowest footprint per user dose would be categorised as Class A.”",
            ],
            "assessment": (
                "Metodyka nazywa klasy zgodnymi z metodyką Ecolabel, która ma zapewniać, że produkty najkorzystniejsze środowiskowo stanowią 10–20% produktów dostępnych na rynku. "
                "W tym samym rozdziale klasa A jest 10% produktów o najniższym śladzie na dawkę w zbiorze, do którego metodyka się stosuje, czyli w produktach L’Oréal. "
                "EU Ecolabel jest w bazie wiedzy znakiem typu I o progu ustalanym dla kategorii, a nie percentylem portfela jednego przedsiębiorcy. "
                "Zestawienie własnego ogona rozkładu z „produktami dostępnymi na rynku” podaje prawdziwą informację o procencie w sposób, który przenosi na literę autorytet Ecolabel (art. 5 ust. 2 pkt 2). "
                "Hasło ogólnej wysokiej efektywności bez tego znaku na produkcie wpada w pkt 4a. W materiale nie wykazano, że kosmetyk Garnier nosi EU Ecolabel."
            ),
            "remediation": (
                "Usunąć z metodyki publikowanej konsumentowi, z kart i z odpowiedzi dla obsługi klienta zdanie o zgodności z metodyką Ecolabel oraz o 10–20% produktów na rynku. "
                "Nie używać słowa Ecolabel przy skali grupy."
            ),
            "claim_title": "SFORMUŁOWANIE NIE DO URATOWANIA",
            "claim": (
                "Zdanie o zgodności skali A–E z metodyką Ecolabel i o 10–20% produktów dostępnych na rynku nie ma bezpiecznej parafrazy. "
                "Usunąć je. EU Ecolabel wolno podać tylko przy produkcie, który ten znak posiada, w zakresie tego znaku."
            ),
            "salvage": False,
        },
        {
            "title": "PRZYPADEK 05: PEF, czternaście czynników i weryfikacja zewnętrzna przy literze",
            "basis": "Podstawa prawna: UPNPR art. 5 ust. 2 pkt 2 i art. 6",
            "risk": "Poziom ryzyka: WYSOKIE",
            "quotes": [
                "„Na ostateczny wynik składa się 14 czynników środowiskowych: od pozyskiwania surowców, poprzez produkcję, transport, aż po korzystanie z produktu końcowego i jego utylizację.”",
                "„aligned with the European Product Environmental Footprint (PEF) guidelines” oraz „Our methodology respects the European Commission’s recommendations on Product Environmental Performance (PEF).”",
                "„While no cosmetic product categories were selected to be part of the official PEF experiment.”",
                "Przy Freshwater Ecotoxicity, Fossil & Mineral Resources Depletion i Toxicity via Environment tabela metodyki podaje „PEF COMPLIANCE: No”.",
                "„data verified by independent auditor Bureau Veritas Certification” oraz „peer-reviewed in January 2020 by four independent experts.”",
            ],
            "assessment": (
                "Wytyczne UCPD 2021 wymagają, aby twierdzenie środowiskowe było prawdziwe, jasne, konkretne, dokładne i jednoznaczne, a dowód dał się zweryfikować. "
                "Powiązanie oceny z PEF jest w tekście spółki prawdziwe co do kierunku metody i nieprawdziwe jako obraz pełnej zgodności. "
                "Kosmetyki nie weszły do oficjalnego eksperymentu PEF. Dla szamponu metodyka powołuje się na grupę cieniową Cosmetics Europe. "
                "Trzy czynniki w tabeli 1 mają zgodność PEF oznaczoną jako „No”: ekotoksyczność wody słodkiej, wyczerpywanie zasobów kopalnych i mineralnych oraz toksyczność przez środowisko. "
                "Akapit polski mówi o 14 czynnikach od surowca do utylizacji i nie podaje tych wyjątków. "
                "Bureau Veritas i recenzja ekspertów, w brzmieniu spółki, dotyczą metody, danych i polityki oznaczania. "
                "Umieszczenie tego zdania na wzorze etykiety obok litery przenosi autorytet audytu na odbiór oceny jako progu środowiskowego. "
                "To transfer atrybutu: opinia o metodzie i weryfikacja danych nie są cechą środowiskową kosmetyku (UPNPR art. 5 ust. 3 pkt 2 w zakresie wyników i zezwoleń oraz art. 6 co do wyjątków PEF)."
            ),
            "remediation": (
                "Przy każdej pozostałej informacji o śladzie, w tym samym bloku, podać datę metodyki, listę czynników liczonych inaczej niż PEF i zakres pracy Bureau Veritas: metody i dane, nie certyfikat produktu. "
                "Zdjąć sformułowania „aligned with PEF” i „respects the recommendations” z miejsc, w których nie ma tego zastrzeżenia."
            ),
            "claim_title": "DOZWOLONE SFORMUŁOWANIE (WZÓR PO KOREKCIE)",
            "claim": (
                "„Ślad policzono według Product Environmental & Social Impact Labelling Methodologies z 25.10.2021. "
                "Kosmetyki nie były kategorią oficjalnego eksperymentu PEF. Ekotoksyczność wody słodkiej, wyczerpywanie zasobów kopalnych i mineralnych oraz toksyczność przez środowisko metodyka liczy wariantem oznaczonym w tabeli jako poza zgodnością PEF. "
                "Bureau Veritas wydał raport o metodach obliczeń i polityce oznaczania. To nie jest certyfikat środowiskowy tego kosmetyku.”"
            ),
            "salvage": True,
        },
        {
            "title": "PRZYPADEK 06: „Responsible plant” i udział energii odnawialnej",
            "basis": "Podstawa prawna: ECGT, załącznik I pkt 4a · UPNPR art. 6",
            "risk": "Poziom ryzyka: WYSOKIE (czarna lista przy haśle „responsible”)",
            "quotes": [
                "„Manufacturing conditions. Made in a responsible plant. Waste recovery: 100%. Renewable energy: 92%.”",
                "„We consider a manufacturing site to be ‘responsible’ when it continuously reduces its CO2 emissions, water consumption and waste production, contributes to the development of local employment (especially for people with disabilities) and gives access to training, irrespective of whether the site belongs to L’Oréal or a subcontractor.”",
                "„renewable energy is energy that is collected from renewable resources, which are naturally replenished on a human timescale, such as sunlight, wind, rain, tides, waves, and geothermal heat.”",
            ],
            "assessment": (
                "Wzorzec oznaczenia stawia przy produkcie zdanie „Made in a responsible plant”. Definicja w metodyce obejmuje ciągłe zmniejszanie emisji CO2, zużycia wody i odpadów oraz zatrudnienie lokalne i szkolenia, także u podwykonawcy. "
                "Jest to opis kierunku, nie próg. „Responsible” przy produkcie jest hasłem ogólnym o charakterze środowiskowym i społecznym. Bez EU Ecolabel lub typu I na ten zakład i na ten zakres wpada w pkt 4a. "
                "Liczba 92% energii odnawialnej jest przykładem na rysunku metodyki, nie pomiarem przypisanym w zebranym materiale do nazwanej partii polskiego SKU. "
                "Definicja wymienia źródła odnawialne i nie podaje przy liczbie, czy chodzi o wytwarzanie na miejscu, umowę na fizyczną dostawę, czy inny instrument, ani jakiego roku i jakiego zakładu dotyczy. "
                "Granica systemu przy procencie jest informacją, bez której liczba zmienia znaczenie (art. 6). "
                "Sprawa Bolt w bazie dotyczy twierdzenia „100% energii odnawialnej” opartego na certyfikatach bez wyjaśnienia. Tego mechanizmu materiał L’Oréal nie opisuje. Zarzut ogranicza się do braku granicy przy liczbie i do hasła „responsible”."
            ),
            "remediation": (
                "Zdjąć „responsible plant” z wzoru oznaczenia. Procent energii publikować z rokiem, nazwą zakładu, nośnikiem (energia elektryczna, ciepło, paliwo) i sposobem nabycia. "
                "Przykład 92% z rysunku metodyki nie może być kopiowany na kartę produktu jako dana tego kosmetyku."
            ),
            "claim_title": "DOZWOLONE SFORMUŁOWANIE (WZÓR PO KOREKCIE)",
            "claim": (
                "„Partia wytworzona w zakładzie [nazwa] w [roku]: udział energii elektrycznej z wytwarzania [słoneczne / wiatrowe / inne, ze wskazaniem umowy] — [X]%. "
                "Liczba dotyczy energii elektrycznej tego zakładu w tym roku. Nie opisuje śladu formuły ani opakowania.”"
            ),
            "salvage": True,
        },
        {
            "title": "PRZYPADEK 07: Odzysk odpadów 100% obejmujący odzysk energii",
            "basis": "Podstawa prawna: UPNPR art. 5 ust. 2 pkt 2 i art. 6",
            "risk": "Poziom ryzyka: ŚREDNIE",
            "quotes": [
                "„Waste recovery: 100%.”",
                "„Part of waste recovery (%): Definition: waste that is either reused, recycled, or recovered with energy generation (such as energy used for urban heating).”",
            ],
            "assessment": (
                "Linia wzoru oznaczenia podaje 100% odzysku odpadów. Definicja w rozdziale o warunkach wytwarzania zalicza do odzysku ponowne użycie, recykling oraz odzysk energii, z przykładem ciepła dla ogrzewania miejskiego. "
                "Odbiorca liczby 100% przy kosmetyku dostaje obraz pełnego zagospodarowania. Udział recyklingu materiałowego i udział spalenia z odzyskiem energii w tej linii nie pada. "
                "Definicja jest w metodyce, a liczba na wzorze etykiety. Odesłanie „learn more” jest osobnym krokiem. "
                "Dla hasła na pierwszym ekranie dopisek na innej stronie nie domyka znaczenia (zasada tego samego kanału, FAQ ECGT i wytyczne UCPD). "
                "To prawda o wskaźniku wewnętrznym podana bez składu wskaźnika (art. 5 ust. 2 pkt 2 i art. 6)."
            ),
            "remediation": (
                "Przy procencie, w tej samej linii, rozdzielić: ponowne użycie, recykling materiałowy, odzysk energii. Podać rok i zakład. "
                "Nie publikować samego „100%”."
            ),
            "claim_title": "DOZWOLONE SFORMUŁOWANIE (WZÓR PO KOREKCIE)",
            "claim": (
                "„Zakład [nazwa], rok [rok]: odpady produkcyjne — recykling materiałowy [A]%, odzysk energii [B]%, inne zagospodarowanie [C]%. "
                "Suma dotyczy odpadów tego zakładu, nie opakowania po produkcie u konsumenta.”"
            ),
            "salvage": True,
        },
        {
            "title": "PRZYPADEK 08: Recyklowalność i recyklat opakowania z wyłączeniem zamknięcia",
            "basis": "Podstawa prawna: ECGT, załącznik I pkt 4b · UPNPR art. 6",
            "risk": "Poziom ryzyka: ŚREDNIE",
            "quotes": [
                "„% of bottle made of recycled material: 74%. Recyclable: yes. Refillable or reusable: no.”",
                "„(3) Excluding closing system. (4) Recycling instructions may vary locally.”",
                "„A packaging or packaging component is recyclable if its successful post-consumer collection, sorting, and recycling is proven to work in practice and at scale.”",
            ],
            "assessment": (
                "Punkt 4b obejmuje twierdzenie o całym produkcie lub całym opakowaniu, gdy dowód dotyczy składnika. "
                "Wzorzec podaje 74% materiału z recyklingu i „yes” przy recyklowalności, a przypis wyłącza system zamknięcia oraz zastrzega, że instrukcja zbiórki zależy od rynku. "
                "Definicja recyklowalności odsyła do formuły Ellen MacArthur Foundation: zbiórka, sortowanie i recykling mają działać w praktyce i na skalę. "
                "Metodyka dopuszcza dostosowanie do rynku lokalnego. W zebranym materiale polskim nie ma zdania, czy „yes” dla butelki obowiązuje przy zbiórce w Polsce i czy dotyczy nakrętki oraz etykiety. "
                "Wzorce z bazy: zarzut wobec kapsułek Tchibo, gdzie warunek recyklingu był w innym języku niż komunikat, oraz istotny pogląd UOKiK w sprawie haseł o butelce „w 100%”, z zastrzeżeniem zakrętki i etykiety poza deklaracją. "
                "Liczby 74% i „yes” są przykładem rysunku, nie pomiarem nazwanej referencji Garnier w Polsce."
            ),
            "remediation": (
                "Przy opakowaniu polskiego SKU podać procent recyklatu z wskazaniem elementu (butelka, bez nakrętki, bez etykiety) oraz czy ten element jest zbierany i przetwarzany w Polsce. "
                "„Recyclable: yes” bez tego zakresu usunąć."
            ),
            "claim_title": "DOZWOLONE SFORMUŁOWANIE (WZÓR PO KOREKCIE)",
            "claim": (
                "„Butelka tego produktu: [X]% tworzywa z recyklingu poużytkowego. Nakrętka i etykieta są poza tą liczbą. "
                "Butelka nadaje się do zbiórki [frakcja, gmina / system] w Polsce. Nakrętki to dotyczy / nie dotyczy.”"
            ),
            "salvage": True,
        },
        {
            "title": "PRZYPADEK 09: Ślad węglowy i wodny wobec średniej własnej kategorii",
            "basis": "Podstawa prawna: UPNPR art. 6 · ISO 14040 i ISO 14044 jako ramy LCA przywołane przez cykl życia w metodyce",
            "risk": "Poziom ryzyka: ŚREDNIE",
            "quotes": [
                "„Carbon footprint 84g per usage dose. 60,3g per 10ml.” Przypis: „Grams of CO2 equivalent (\"shampoo\" category average = 100g).”",
                "„Water footprint 6,1 per usage dose.” Przypis: „Water index including water quality + water scarcity (\"shampoo\" category average = 7,7).”",
                "„so as to allow consumers compare products offering the same type of cosmetic benefit and be able to select the product with the lowest environmental footprint.”",
            ],
            "assessment": (
                "Liczba CO2 przy claimie komercyjnym wymaga metodyki, granic systemu, jednostki funkcjonalnej i scenariusza. "
                "Metodyka podaje jednostkę: dawka użycia oraz objętość 10 ml, a etapy: od wydobycia składników formuły i materiałów opakowania przez produkcję, dystrybucję i użycie do końca życia. "
                "To jest rama, której sam akapit polski o „14 czynnikach” nie powtarza przy liczbie. "
                "Średnia 100 g i 7,7 przy śladzie wody jest średnią kategorii „shampoo” w logice tej samej skali, czyli kategorii portfela, nie średnią rynku. "
                "Na rysunku średnia stoi przy gramach jako punkt odniesienia. Odbiorca porównuje 84 g ze 100 g i może uznać różnicę za przewagę nad kategorią rynkową. "
                "Część czynników wagowych jest poza zgodnością PEF (przypadek 05). Streszczenia z celem, granicami, datą danych i dostępem do skrótu przy liczbie na karcie w materiale polskim nie ma. "
                "Wytyczne UCPD 2021: powołanie LCA wymaga metod uznanych dla danego wyrobu i weryfikacji strony trzeciej, a założenia mają być dla odbiorcy dostępne."
            ),
            "remediation": (
                "Jeżeli liczba zostaje, przenieść na kartę produktu: wartość, jednostkę, granice, datę, informację o czynnikach poza PEF i zdanie, że średnia dotyczy produktów L’Oréal w tej kategorii. "
                "Przykładowe 84 g i 100 g z rysunku szamponu nie publikować jako wyniku innego kosmetyku."
            ),
            "claim_title": "DOZWOLONE SFORMUŁOWANIE (WZÓR PO KOREKCIE)",
            "claim": (
                "„Ślad węglowy jednej dawki użycia: [X] g CO2e. Granice: surowce formuły i opakowania, produkcja, transport, użycie, koniec życia opakowania. Metodyka: PIL z 25.10.2021, ze wskazaniem czynników poza zgodnością PEF. "
                "Średnia [Y] g CO2e na dawkę dotyczy produktów L’Oréal w kategorii [nazwa], stan na [data]. To nie jest średnia rynku i nie jest EU Ecolabel.”"
            ),
            "salvage": True,
        },
        {
            "title": "PRZYPADEK 10: Wpływ społeczny z liczby dostawców",
            "basis": "Podstawa prawna: UPNPR art. 5 ust. 3 pkt 2 i art. 6 · ECGT, załącznik I pkt 4b",
            "risk": "Poziom ryzyka: ŚREDNIE",
            "quotes": [
                "„A oznacza kosmetyk najlepszy w swojej klasie pod względem wpływu na środowisko i społeczności.”",
                "„The ingredients and components of this product are sourced from suppliers committed to respect the fundamental principles of the UN on labor standards. 3 suppliers committed to the social inclusion of people from vulnerable communities contributed to making this product.”",
            ],
            "assessment": (
                "Blok społeczny wzoru oznaczenia stwierdza, że składniki i komponenty pochodzą od dostawców zobowiązanych do zasad pracy ONZ oraz że trzech dostawców zaangażowanych we włączenie osób z grup wrażliwych przyczyniło się do powstania produktu. "
                "Jest to cecha polityk dostawców i liczba podmiotów, nie pomiar skutku społecznego formuły. "
                "Litera A w tekście polskim obejmuje „wpływ na społeczności” łącznie ze środowiskiem. Litera w metodyce jest klasyfikacją śladu środowiskowego. "
                "Przeniesienie zobowiązania dostawców i liczby „3” na ocenę produktu jest transferem atrybutu (art. 5 ust. 3 pkt 2: cechy i spodziewane wyniki). "
                "Punkt 4b obejmuje twierdzenie o całym produkcie, gdy informacja dotyczy wycinka łańcucha. Trzech dostawców nie opisuje społeczności, na które produkt oddziałuje."
            ),
            "remediation": (
                "Oddzielić moduł społeczny od śladu. Podać, ilu dostawców jakiej kategorii składników objęto jakim standardem, w jakim roku i jakim audytem. "
                "Nie wiązać tej liczby z literą oceny środowiskowej."
            ),
            "claim_title": "DOZWOLONE SFORMUŁOWANIE (WZÓR PO KOREKCIE)",
            "claim": (
                "„Dostawcy składników tego produktu, objęci programem w [roku]: [n] podmiotów z zobowiązaniem do zasad pracy ONZ, weryfikacja: [nazwa i data]. "
                "To informacja o dostawcach. Nie jest oceną środowiskową kosmetyku.”"
            ),
            "salvage": True,
        },
        {
            "title": "PRZYPADEK 11: 96% produktów o poprawionym profilu",
            "basis": "Podstawa prawna: ECGT, załącznik I pkt 4b · UPNPR art. 5 ust. 2 pkt 2",
            "risk": "Poziom ryzyka: ŚREDNIE",
            "quotes": [
                "„Since 2017, all new or renovated L'Oréal products are evaluated with SPOT. By 2020, 96% of them had an an improved profile.”",
                "„we will improve its environmental or social profile against at least one of the following criteria: the new formula reduces the environmental footprint; the new formula uses raw materials that are sustainably sourced or that or derived from Green Chemistry; the new packaging has an improved environmental profile; the new product a positive social impact.”",
            ],
            "assessment": (
                "96% dotyczy nowych lub odnowionych produktów ocenionych w SPOT, które do 2020 r. miały poprawiony profil. "
                "Przypis metodyki definiuje poprawę jako zmianę wobec co najmniej jednego kryterium: niższy ślad formuły, surowiec z deklarowanego zrównoważonego pozyskania lub z zielonej chemii, lepszy profil opakowania albo dodatni skutek społeczny. "
                "Spełnienie jednego z tych punktów, w tym kryterium społecznego, wystarcza do zaliczenia produktu. "
                "Zdanie o 96% stoi we wstępie metodyki oznaczenia, które konsumentowi przedstawia wpływ środowiskowy. "
                "Odbiorca czyta odsetek jako poprawę profilu środowiskowego portfolio. Dowód spółki jest węższy (pkt 4b i art. 5 ust. 2 pkt 2). "
                "Wzorzec Zary w bazie: certyfikat i procent na karcie sugerujące, że niemal cały model spełnia standard, przy zakresie węższym niż komunikat."
            ),
            "remediation": (
                "Przy 96% powtórzyć w tym samym akapicie: licznik (nowe i odnowione, z wyłączeniem niedawnych akwizycji, jeśli tak stanowi przypis), rok bazowy i zasadę „co najmniej jedno kryterium”, z rozbiciem na kryteria środowiskowe i społeczne. "
                "Bez tego rozbicia odsetka nie publikować obok skali produktu."
            ),
            "claim_title": "DOZWOLONE SFORMUŁOWANIE (WZÓR PO KOREKCIE)",
            "claim": (
                "„W 2020 r. 96% nowych lub odnowionych produktów L’Oréal ocenionych w SPOT spełniało co najmniej jedno z kryteriów: niższy ślad formuły, surowiec z programu pozyskania spółki, zmiana opakowania albo skutek społeczny. "
                "Odsetek nie oznacza, że 96% produktów ma niższy ślad środowiskowy we wszystkich kategoriach wpływu.”"
            ),
            "salvage": True,
        },
    ]

    for case in cases:
        hpar = doc.add_paragraph()
        hpar.paragraph_format.space_before = Pt(12)
        hpar.paragraph_format.space_after = Pt(2)
        hpar.paragraph_format.keep_with_next = True
        _set_run(hpar.add_run(case["title"]), 12, True, False, NAVY)
        meta_line = doc.add_paragraph()
        meta_line.paragraph_format.space_before = Pt(0)
        meta_line.paragraph_format.space_after = Pt(4)
        _set_run(meta_line.add_run(case["basis"] + "  |  " + case["risk"]), 9, True, False, DANGER if "BARDZO" in case["risk"] or "WYSOKIE" in case["risk"] else WARNING)
        add_p("Cytat źródłowy:", size=10.5, bold=True, after=2)
        for q in case["quotes"]:
            add_p(q, size=10.5, italic=True, after=3)
        add_p("Ocena prawna:", size=10.5, bold=True, before=4, after=2)
        add_p(case["assessment"])
        add_p("Kierunek naprawy:", size=10.5, bold=True, before=2, after=2)
        add_p(case["remediation"])
        bg = "ECFDF5" if case["salvage"] else "F8FAFC"
        border = "059669" if case["salvage"] else "94A3B8"
        title_color = RGBColor(4, 120, 87) if case["salvage"] else MUTED
        callout(case["claim_title"], [case["claim"]], bg, border, title_color)

    add_h("4. Porównanie twierdzeń L’ORÉAL POLSKA z rzeczywistym stanem prawno-środowiskowym")
    add_p(
        "Sekcja dotyczy porównań, które materiał sam buduje: klasa produktu, udział w rynku przy metodyce Ecolabel, zgodność z PEF oraz średnia kategorii przy gramach CO2e.",
        after=6,
    )
    comp = [
        (
            "Klasa oceny. „A oznacza kosmetyk najlepszy w swojej klasie pod względem wpływu na środowisko i społeczności.” Raport: „A is the most favourable profile for the planet”, przy porównaniu z produktami grupy.",
            "Klasa A w metodyce to 10% produktów L’Oréal o najniższym śladzie na dawkę w kategorii portfela. Klasa E to 10% o śladzie najwyższym. Litera jest śladem środowiskowym. Skutek społeczny jest osobnym blokiem o dostawcach. UZNK art. 16 ust. 3 wymaga porównania sprawdzalnego wobec towarów tej samej potrzeby. Zbiór własny tego warunku nie spełnia.",
        ),
        (
            "Udział rynkowy i Ecolabel. Klasy „compliant with the Ecolabel methodology”, która ma zapewniać, że produkty najkorzystniejsze stanowią 10–20% produktów dostępnych na rynku.",
            "Zbiór, z którego metodyka wycina 10%, to produkty L’Oréal. EU Ecolabel (rozporządzenie 66/2010) jest oznakowaniem typu I o progu kategorii, nie percentylem jednego portfela. Hasło wysokiej efektywności bez tego znaku na produkcie jest objęte ECGT pkt 4a.",
        ),
        (
            "PEF. Metodyka „aligned with the European Product Environmental Footprint (PEF) guidelines”; raport: metodyka „respects” zalecenia Komisji dotyczące PEF. Polski akapit: 14 czynników od surowca do utylizacji.",
            "Kosmetyki nie weszły do oficjalnego eksperymentu PEF. Tabela metodyki oznacza „PEF COMPLIANCE: No” przy ekotoksyczności wody słodkiej, wyczerpywaniu zasobów kopalnych i mineralnych oraz toksyczności przez środowisko. Bureau Veritas weryfikuje metody i dane, nie nadaje znaku typu I.",
        ),
        (
            "Średnia przy śladzie. Przykład: 84 g CO2e na dawkę, średnia kategorii „shampoo” = 100 g; indeks wody 6,1 wobec średniej 7,7.",
            "Średnia jest średnią kategorii w tej samej skali, czyli portfela grupy. Jednostka (dawka użycia) i etapy cyklu życia są w PDF metodyki. Na wzorze oznaczenia średnia stoi przy gramach jako punkt odniesienia, bez zdania, że jest to średnia wewnętrzna.",
        ),
    ]
    ctable = doc.add_table(rows=1 + len(comp), cols=2)
    ctable.alignment = WD_TABLE_ALIGNMENT.CENTER
    ctable.autofit = False
    ctable.style = "Table Grid"
    cw = (Cm(8.8), Cm(8.8))
    for c, htxt in enumerate(["Twierdzenie podmiotu", "Rzeczywisty stan prawno-środowiskowy"]):
        cell = ctable.rows[0].cells[c]
        cell.width = cw[c]
        shade(cell, "08181F")
        margins(cell, 50, 50, 60, 60)
        par = cell.paragraphs[0]
        par.paragraph_format.space_before = par.paragraph_format.space_after = Pt(1)
        _set_run(par.add_run(htxt), 8, True, False, WHITE)
    for r, (left, right) in enumerate(comp):
        bg = "FFFFFF" if r % 2 == 0 else "F8FAFC"
        for c, val in enumerate((left, right)):
            cell = ctable.rows[r + 1].cells[c]
            cell.width = cw[c]
            shade(cell, bg)
            margins(cell, 50, 50, 60, 60)
            borders(cell)
            par = cell.paragraphs[0]
            par.paragraph_format.space_before = par.paragraph_format.space_after = Pt(1)
            _set_run(par.add_run(val), 8, False, False, DARK)

    add_h("5. Ryzyka prawne i finansowe")
    risks = [
        (
            "1. Postępowanie Prezesa UOKiK",
            "Komunikat z 16 września 2026 r. stwierdza zarzut naruszania zbiorowych interesów konsumentów w związku z komunikacją środowiskową kosmetyków Garnier. Praktyka skali na stronie ustała w czerwcu 2025 r., w toku postępowania. Zaniechanie nie zastępuje decyzji. Art. 24 ustawy o ochronie konkurencji i konsumentów obejmuje nieuczciwe praktyki rynkowe i obowiązek informacji rzetelnej, prawdziwej i pełnej. Art. 106 ust. 1 pkt 4 dopuszcza karę do 10% obrotu za naruszenie art. 24, także nieumyślne. W zebranym materiale kary nie nałożono.",
        ),
        (
            "2. Roszczenia konkurentów",
            "„Najlepszy w swojej klasie” pozwala rozpoznać kosmetyki tej samej funkcji. UZNK art. 16 ust. 1 pkt 2 i ust. 3 czynią z reklamy porównawczej czyn nieuczciwej konkurencji, gdy wprowadza klienta w błąd albo nie porównuje cech w sposób rzetelny i sprawdzalny. Konkurent może żądać zaniechania i naprawienia szkody. Art. 10 UZNK obejmuje także oznaczenie towaru mogące wprowadzać w błąd co do jakości.",
        ),
        (
            "3. ECGT na kanale internetowym",
            "Od 27 września 2026 r. zakazy załącznika, w tym pkt 2a, 4a i 4b, są miarą oceny twierdzeń B2C. Porozumienie CPC o starych zapasach nie daje ulgi dla strony internetowej: korektę online organy traktują jako natychmiastową. Tytuł „Ekologiczne kosmetyki” i meta o ocenie wpływu były opublikowane 11 października 2026 r. Polska nowela (druk Senatu 825, wejście w życie 27 września 2027 r. po zakończeniu prac) nie wstrzymuje tej oceny. Do czasu ogłoszenia noweli Prezes UOKiK opiera zarzuty na obecnym UPNPR.",
        ),
        (
            "4. Sprawozdawczość i dane dla kontrahentów",
            "Ten sam odsetek „improved profile”, litera A i średnia kategorii, przeniesione do raportu zrównoważonego rozwoju albo do ankiety kontrahenta objętego CSRD, powielają względny percentyl jako informację o efektywności środowiskowej. Odbiorca danych scope 3 dostaje wtedy pozycję w portfelu grupy w miejscu progu bezwzględnego.",
        ),
    ]
    for title, body in risks:
        add_p("• " + title, size=11, bold=True, before=8, after=2)
        add_p(body)

    add_h("6. Rekomendacje naprawcze")
    stages = [
        (
            "Etap 1: natychmiastowy pull-off — 11 października 2026 r.",
            "Zdjąć z garnier.pl tytuł „Ekologiczne kosmetyki”, meta „oceniane pod względem wpływu na środowisko” i każde pozostałe „Wybieraj zrównoważone produkty” oraz literę A–E, jeżeli występuje na karcie, w marketplace, w newsletterze lub w pliku do pobrania. Nie wznawiać skali. Zlecić ponowne zindeksowanie adresów kategorii.",
            "FFF1F2", "E11D48", DANGER,
        ),
        (
            "Etap 2: korekta kart i odpowiedzi w postępowaniu — do 18 października 2026 r.",
            "Zestawić miejsca, w których skala lub hasło ogólne jeszcze wraca (wyszukiwarka, porównywarki, ulotki, social). W piśmie do Prezesa UOKiK opisać datę zaprzestania modułu i listę adresów poprawionych po 11 października 2026 r. Usunąć z publicznego skrótu metodyki zdanie o zgodności z Ecolabel i o 10–20% rynku.",
            "FFFBEB", "D97706", WARNING,
        ),
        (
            "Etap 3: metodyka przy liczbie, w tym samym widoku — do 10 listopada 2026 r.",
            "Jeżeli spółka zostawia ślad liczbowy: na karcie produktu dawka użycia, granice cyklu życia, data metodyki, czynniki poza zgodnością PEF, zbiór odniesienia „produkty L’Oréal w kategorii [nazwa]” oraz rozdzielenie modułu społecznego. Procent energii, odzysk odpadów i recyklat opakowania — z rokiem, zakładem i wyłączeniami (nakrętka, odzysk energii) w tej samej linii co liczba.",
            "F0F9FF", "0284C7", ACCENT,
        ),
        (
            "Etap 4: safe harbor i szkolenie — tryb stały",
            "Zatwierdzić zamknięty katalog zdań: surowiec i procent na nazwanym elemencie opakowania, ślad na dawkę z granicami, program dostawców z rokiem i audytem. Zakaz: litera A–E jako znak, „ekologiczny”, „zrównoważony”, „responsible”, „Ecolabel”, „najlepszy w klasie”, „profil najkorzystniejszy dla planety”. Szkolenie marketingu Garnier i osób aktualizujących title oraz meta.",
            "ECFDF5", "059669", RGBColor(4, 120, 87),
        ),
    ]
    for title, body, bg, border, color in stages:
        callout(title, [body], bg, border, color)

    add_h("7. Wnioski")
    add_p(
        "Skala A–E Garnier jest własnym oznakowaniem grupy i pozycją w rozkładzie portfela L’Oréal, a komunikat dla konsumenta w Polsce nadaje literze A znaczenie najlepszego w klasie skutku środowiskowego i społecznego. "
        "W tej formie materiału nie publikować. Zarzut Prezesa UOKiK z 16 września 2026 r. pozostaje zarzutem: zebrane źródła nie zawierają decyzji ani wymierzonej kary."
    )

    out = ROOT / "Raport_Audytu_Greenwashingu_LOreal.docx"
    doc.save(out)
    return out


if __name__ == "__main__":
    print(build_material())
    print(build_report())
